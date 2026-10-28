import re
from io import BytesIO
from collections import Counter

STOPWORDS = set('a an the and or but if then than to of in on for with by at from as is are was were be been being it this that these those your you we our their they he she his her i me my us will can could should would may might have has had do does did not into over under about across through using used use such also more most very'.split())
ACTION_VERBS = set('built created developed designed implemented improved managed led analyzed analysed automated delivered launched reduced increased optimized optimised collaborated coordinated tested maintained deployed integrated supported resolved trained documented proposed streamlined achieved conducted engineered configured'.split())
SECTIONS = {
    'Contact details': [r'\b(email|phone|mobile|linkedin|github|portfolio)\b'],
    'Summary / objective': [r'\b(summary|objective|profile)\b'],
    'Skills': [r'\bskills\b'],
    'Experience': [r'\b(experience|employment|work history|internship)\b'],
    'Projects': [r'\bprojects?\b'],
    'Education': [r'\beducation\b', r'\bacademic\b'],
    'Certifications': [r'\b(certifications?|certificates|licenses)\b']
}

def clean_text(value):
    return value.replace('\x00', '').strip() if isinstance(value, str) else ''

def generate_resume(data):
    val = lambda key: clean_text(data.get(key, ''))
    name = val('name') or 'Your Name'
    contact = ' | '.join(x for x in [val('email'), val('phone'), val('location'), val('linkedin'), val('github')] if x)
    lines = [name.upper()]
    if val('headline'): lines.append(val('headline'))
    if contact: lines.append(contact)
    for heading, key in [('Professional Summary','summary'), ('Technical Skills','skills'), ('Experience','experience'), ('Projects','projects'), ('Education','education'), ('Certifications','certifications')]:
        if val(key): lines.extend(['', heading.upper(), val(key)])
    return '\n'.join(lines).strip()

def tokens(text):
    return re.findall(r"[a-zA-Z][a-zA-Z0-9+#.\-]{1,}", text.lower())

def terms(text):
    return Counter(t for t in tokens(text) if t not in STOPWORDS and len(t) > 2)

def analyze_resume(resume_text, job_description=''):
    words = tokens(resume_text)
    resume_terms = set(terms(resume_text))
    jd_terms = terms(job_description)
    matched = sorted(t for t in jd_terms if t in resume_terms)
    missing = sorted((t for t in jd_terms if t not in resume_terms), key=lambda t: (-jd_terms[t], t))[:15]
    keyword_score = round(100 * len(matched) / len(jd_terms)) if jd_terms else None
    section_status = {name: any(re.search(p, resume_text, re.I) for p in patterns) for name, patterns in SECTIONS.items()}
    checks = [
        (bool(re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', resume_text)), 10, 'Add a professional email address.'),
        (bool(re.search(r'(\+?\d[\d\s().-]{7,}\d)', resume_text)), 8, 'Add a reachable phone number.'),
        (section_status['Summary / objective'], 10, 'Add a short professional summary or objective.'),
        (section_status['Skills'], 12, 'Include a clearly labelled Skills section.'),
        (section_status['Education'], 10, 'Include your education and institution.'),
        (section_status['Projects'] or section_status['Experience'], 14, 'Add relevant projects, work experience, or internships.'),
        (section_status['Certifications'], 5, 'Add relevant certifications if you have them.'),
        (len(words) >= 120, 10, 'Add more relevant detail if the resume is very short.'),
        (len(words) <= 900, 5, 'Keep the resume focused and remove repeated details.'),
        (bool(set(tokens(resume_text)) & ACTION_VERBS), 8, 'Use action verbs such as developed, implemented, or improved.'),
        (bool(re.search(r'\b\d+%|\b\d+\+?\s*(users|customers|projects|hours|days|weeks|months|years|records|requests)\b', resume_text, re.I)), 8, 'Add measurable outcomes where accurate.')
    ]
    score = min(100, sum(points for passed, points, _ in checks if passed))
    suggestions = [msg for passed, _, msg in checks if not passed]
    if job_description and missing:
        suggestions.insert(0, 'Tailor your resume with relevant missing terms only when they truthfully reflect your skills or experience.')
    if not job_description:
        suggestions.insert(0, 'Paste a target job description to calculate keyword overlap for that role.')
    if not suggestions:
        suggestions.append('Your resume covers the main structure checks. Review all claims and tailor it to each role.')
    return {
        'score': score, 'score_label': 'Resume completeness score', 'keyword_score': keyword_score,
        'word_count': len(words), 'unique_words': len(set(words)),
        'sections': [{'name': k, 'present': v} for k, v in section_status.items()],
        'matched_keywords': matched[:25], 'missing_keywords': missing,
        'suggestions': suggestions[:10],
        'disclaimer': 'This is a heuristic guide, not a real ATS vendor score or a guarantee of hiring outcomes. Different employers and systems evaluate resumes differently.'
    }

def extract_text_from_upload(file_storage, extension):
    raw = file_storage.read()
    if extension == 'txt':
        return raw.decode('utf-8-sig', errors='replace')
    if extension == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(raw))
        return '\n'.join(page.extract_text() or '' for page in reader.pages)
    if extension == 'docx':
        from docx import Document
        doc = Document(BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(' | '.join(cell.text.strip() for cell in row.cells))
        return '\n'.join(parts)
    raise ValueError('Unsupported file type')

def render_analysis_pdf(analysis):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from xml.sax.saxutils import escape

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=.65*inch, rightMargin=.65*inch, topMargin=.55*inch, bottomMargin=.55*inch, title='Resume Analysis Report')
    styles = getSampleStyleSheet()
    title = ParagraphStyle('ReportTitle', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=20, textColor=colors.HexColor('#16243a'), spaceAfter=12)
    heading = ParagraphStyle('ReportHeading', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=11, textColor=colors.HexColor('#5267e8'), spaceBefore=10, spaceAfter=5)
    body = ParagraphStyle('ReportBody', parent=styles['BodyText'], fontName='Helvetica', fontSize=9, leading=13, textColor=colors.HexColor('#344054'))
    story = [Paragraph('Resume Analysis Report', title), Paragraph(f"Overall completeness score: <b>{analysis['score']}/100</b>", body), Spacer(1, 6)]

    if analysis['keyword_score'] is not None:
        story.append(Paragraph(f"Job keyword overlap: <b>{analysis['keyword_score']}%</b>", body))
    story.append(Paragraph(f"Words: {analysis['word_count']}  |  Unique words: {analysis['unique_words']}", body))

    story += [Paragraph('Section Checklist', heading)]
    section_rows = [['Section', 'Status']] + [[s['name'], 'Present' if s['present'] else 'Missing'] for s in analysis['sections']]
    section_table = Table(section_rows, colWidths=[4.5*inch, 1.4*inch])
    section_table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#5267e8')), ('TEXTCOLOR',(0,0),(-1,0),colors.white), ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'), ('GRID',(0,0),(-1,-1),0.5,colors.HexColor('#d9dfeb')), ('FONTNAME',(0,1),(0,-1),'Helvetica'), ('FONTSIZE',(0,0),(-1,-1),9), ('BACKGROUND',(0,1),(-1,-1),colors.HexColor('#f8faff')), ('TOPPADDING',(0,0),(-1,-1),6), ('BOTTOMPADDING',(0,0),(-1,-1),6)]))
    story.append(section_table)

    story += [Paragraph('Matched Keywords', heading), Paragraph(escape(', '.join(analysis['matched_keywords']) or 'None found'), body), Paragraph('Missing Skills & Keywords', heading), Paragraph(escape(', '.join(analysis['missing_keywords']) or 'No missing keywords identified'), body), Paragraph('Suggestions', heading)]
    for suggestion in analysis['suggestions']:
        story.append(Paragraph('• ' + escape(suggestion), body))
    story += [Spacer(1, 10), Paragraph(escape(analysis['disclaimer']), body)]
    doc.build(story)
    return buffer.getvalue()

def render_resume_pdf(text, name='Resume'):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from xml.sax.saxutils import escape
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=.65*inch, leftMargin=.65*inch, topMargin=.55*inch, bottomMargin=.55*inch, title=f'{name} - Resume')
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle('ResumeName', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=17, leading=20, alignment=TA_CENTER, spaceAfter=8)
    heading_style = ParagraphStyle('ResumeHeading', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=10, leading=13, spaceBefore=9, spaceAfter=3)
    body_style = ParagraphStyle('ResumeBody', parent=styles['BodyText'], fontName='Helvetica', fontSize=9, leading=12, spaceAfter=3)
    lines = text.splitlines()
    first = next((i for i, line in enumerate(lines) if line.strip()), 0)
    story = []
    for i, line in enumerate(lines):
        line = line.strip()
        if not line:
            story.append(Spacer(1, 3)); continue
        safe = escape(line).replace('  ', '&nbsp;&nbsp;')
        if i == first: story.append(Paragraph(safe, name_style))
        elif line.upper() == line and len(line) < 60 and not line.endswith('.'): story.append(Paragraph(safe, heading_style))
        else: story.append(Paragraph(safe, body_style))
    doc.build(story)
    return buffer.getvalue()
